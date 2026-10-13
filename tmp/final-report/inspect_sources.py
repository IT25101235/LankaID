from pathlib import Path
from docx import Document
import json, zipfile, hashlib
from pypdf import PdfReader
root=Path(r'D:\SLIIT\Year 02\Semester 01\SE\Project')
out=Path(__file__).parent
sources={'nusla':root/'Final Report/2026-Y2-S1-MTR-02_nusla.docx','yasiru':root/'Final Report/SE2030-Final-Report-2026-Y2-S1-MTR-02-yasiru.docx','specification':root/'SE2030_Group Project Specification.pdf','design':root/'Design Document/Uploading Document/2026-Y2-S1-MTR-02.pdf'}
for name,path in sources.items():
 if path.suffix=='.pdf':
  doc=PdfReader(path).pages
  text='\n'.join(f'\n--- PAGE {i+1} ---\n'+p.extract_text() for i,p in enumerate(doc))
  print(name,'pages',len(doc))
 else:
  doc=Document(path)
  text='\n'.join(f'P{i} [{p.style.name}] {p.text}' for i,p in enumerate(doc.paragraphs))
  for i,t in enumerate(doc.tables):
   text+=f'\nTABLE {i}\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows)
  info={'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sections':[{'width':s.page_width.inches,'height':s.page_height.inches,'margins':[s.top_margin.inches,s.bottom_margin.inches,s.left_margin.inches,s.right_margin.inches]} for s in doc.sections], 'styles':{s.name:{'font':s.font.name,'size':s.font.size.pt if s.font.size else None} for s in doc.styles if s.name in ['Normal','Title','Heading 1','Heading 2','Heading 3']}}
  (out/f'{name}_audit.json').write_text(json.dumps(info,indent=2),encoding='utf8')
  z=zipfile.ZipFile(path)
  media=out/f'{name}_media';media.mkdir(exist_ok=True)
  for n in z.namelist():
   if n.startswith('word/media/'):(media/Path(n).name).write_bytes(z.read(n))
  print(name,'paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'images',len(list(media.iterdir())))
 (out/f'{name}.txt').write_text(text,encoding='utf8')

