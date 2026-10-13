from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import zipfile, hashlib
ROOT=Path.cwd(); TMP=Path(__file__).parent
SOURCE=Path(r'D:\SLIIT\Year 02\Semester 01\SE\Project\Final Report\2026-Y2-S1-MTR-02_nusla.docx')
FINAL=ROOT/'output/final-report/2026-Y2-S1-MTR-02_nusla_final.docx'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()=='38d106265310c52e50532c460cbe0a89b24e142a1c6bb6d94cd3871191a5fc04'
doc=Document(SOURCE); body=doc._element.body
sects=[deepcopy(s._sectPr) for s in doc.sections]
sdt=deepcopy(next(e for e in body if e.tag==qn('w:sdt')))
for instruction in sdt.findall('.//'+qn('w:instrText')):instruction.text=' TOC \\o "1-1" \\h \\z '
for e in list(body)[16:]:body.remove(e)
body.append(deepcopy(sects[-1]))
if 'Title' not in doc.styles:doc.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',12),('Title',14)]:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.bold=True
 s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
p=doc.paragraphs[2];p.style=doc.styles['Title'];p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for i in [2,3,4,5]:
 for r in doc.paragraphs[i].runs:r.font.color.rgb=RGBColor(0,0,0)
def para(text='',size=None):
 p=doc.add_paragraph(text)
 if size:
  for r in p.runs:r.font.size=Pt(size)
 return p
def h(text,level=2):return doc.add_paragraph(text,style=f'Heading {level}')
def page(text=None):
 p=doc.add_paragraph();p.paragraph_format.page_break_before=True;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
 if text:h(text)
def section_end(index):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1);p._p.get_or_add_pPr().append(deepcopy(sects[index]))
def table(headers,rows,widths,size=10):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,w,s in zip(t.rows[0].cells,widths,headers):c.width=Inches(w);c.text=s
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 for row in rows:
  for c,w,text in zip(t.add_row().cells,widths,row):c.width=Inches(w);c.text=str(text)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement(f'w:{edge}');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for edge,val in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
    e=OxmlElement(f'w:{edge}');e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');m.append(e)
   pr.append(m);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7' if ri==0 else ('F5F5F5' if ri%2==0 else 'FFFFFF'));pr.append(shade)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.02
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0);r.bold=(ri==0)
 return t
def figure(file,caption,width=6.2,height=None):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
 shape=p.add_run().add_picture(str(file),width=Inches(width))
 if height and shape.height>Inches(height):
  shape.width=int(shape.width*Inches(height)/shape.height);shape.height=Inches(height)
 shape._inline.docPr.set('descr',caption)
 p=para(caption,9.5);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(7)
media=lambda n:TMP/f'yasiru_media/image{n}.png'
body.insert(len(body)-1,sdt)
h('Figures',2)
table(['Figure','Title','Page'],[
('1','Citizen and application-officer use cases','5'),
('2','Public visitor and operational use cases','6'),
('3','Accounts, applications and review relationships','7'),
('4','Appointments and supporting relationships','8'),
('5','Citizen and staff login','9'),
('6','Application tracking and next action','9'),
('7','Biometric appointment selection','10'),
('8','QR appointment pass','10'),
('9','NIC review queue','12'),
('10','Evidence and officer decisions','12'),
('11','Officer slot schedule','13'),
('12','QR and manual check-in','13')],[.55,5.35,.6],10)
section_end(1)
h('Introduction',1)
h('Project overview')
para('LankaID is a Java web application for managing National Identity Card (NIC) and passport services in one portal. Citizens can register, submit a new application or request renewal or replacement, upload supporting evidence, follow officer feedback, pay a simulated fee and book a biometric appointment. Staff review applications, record centre attendance and biometric completion, prepare an issuance register, and manage delivery or collection.')
para('The five-member team developed the prototype for SE2030 Software Engineering. Requirements analysis, UML modelling, iterative development, object-oriented implementation and testing support a connected citizen-service workflow [1], [2]. Payments, biometric device capture and physical document production remain simulations; the portal records their workflow outcomes.')
h('Problem and objectives')
para('The project addresses fragmented application handling: paper-based evidence, repeated visits, separate staff records and limited visibility of application progress. A shared portal can reduce repeated data entry and expose the next action to the applicant and responsible officer. These benefits are design goals; processing-time savings have not been measured in a government deployment.')
para('The objectives are to provide online applications with saved drafts; make review decisions traceable; enforce appointment capacity and change rules; present progress and notifications; support assisted walk-in services; and give administrators and the director operational oversight. The implementation connects these goals through one database and controlled status changes.')
h('Users and stakeholders')
table(['User or stakeholder','Purpose and access'],[('Citizen','Applies, corrects evidence, pays, books, tracks and chooses delivery or collection.'),('NIC and passport officers','Review their assigned document type, manage appointment slots, check in citizens and record biometric completion.'),('CSR and walk-in visitor','CSR handles support, assisted visits and collection. A visitor uses the walk-in route with mobile verification.'),('Director and administrator','Director handles referrals, reports and issuance registers. Administrator manages accounts, centres, fees, notices and maintenance.')],[1.65,4.85],10)
h('Scope and boundaries')
para('The prototype covers registration through recorded handover for new, renewed and replaced documents. It assumes internet access, an English-language browser interface and configured SQL Server. It does not connect to live government identity records, a certified payment gateway, biometric hardware, courier tracking or physical printing. Uploaded files use local server storage.')
section_end(2)
h('Requirements',1)
h('Functional requirements for citizen services')
para('The team organised requirements around the personas and product backlog in the design document [2]. Each identifier links a user need to a verifiable outcome. The requirements describe the prototype scope rather than statutory government procedures.')
table(['ID','Requirement and acceptance outcome','Primary user'],[
('FR01','Verify email and a Sri Lankan mobile number before account creation. Incorrect, expired or over-attempted codes must be rejected.','Citizen'),
('FR02','Authenticate users, maintain a session and restrict protected routes by role and ownership. Logout ends access; password recovery uses email.','Account users'),
('FR03','Create, save, reopen, edit or discard a draft NIC/passport application. Submission validates required fields and produces an application reference.','Citizen'),
('FR04','Upload supporting evidence, preview authorised files and replace flagged evidence. The application service rejects files larger than 5 MB.','Citizen'),
('FR05','Display status, history and the next action. A correction request identifies what must change before resubmission.','Citizen'),
('FR06','Link an identity record in the prototype database and create eligible NIC/passport renewal or lost-document replacement requests.','Citizen'),
('FR07','Show application and optional delivery fees. Validate simulated card input and record a payment outcome and transaction reference.','Citizen'),
('FR08','Book a published biometric slot after workflow prerequisites are met. Confirmation must not exceed capacity and must produce a QR pass.','Citizen'),
('FR09','Cancel or reschedule an eligible booking. Confirmed biometric bookings can change at least 72 hours before the visit; missed bookings can be rescheduled after their date.','Citizen'),
('FR10','Select home delivery or centre collection, view delivery stages and confirm receipt or staff-assisted collection.','Citizen / CSR')],[.48,4.95,1.07],10)
h('Citizen service flow')
para('The normal path is registration, validated submission, officer review, approval, simulated payment, appointment booking, centre check-in, recorded biometric completion and issuance/handover. Correction returns the application to the citizen; rejection records a reason. Downstream actions check prerequisites so a page link cannot bypass the workflow.')
page('Functional requirements for staff and supporting services')
table(['ID','Requirement and acceptance outcome','Primary user'],[
('FR11','Keep NIC/passport review queues separate. Officers inspect details and evidence, approve, reject with a reason, request correction or refer a case to the director.','NIC/passport officer'),
('FR12','Publish and maintain centre/date/time slots with capacity. Exclude closed/full slots and prevent concurrent requests taking the same final place.','Officer'),
('FR13','Locate an appointment by webcam QR scan or reference. Permit check-in on the booked date; record attendance and simulated biometric completion with staff identity.','Officer'),
('FR14','Verify a visitor mobile by SMS, book an assisted walk-in visit, make a simulated visit payment, recover a pass and record the CSR outcome.','Visitor / CSR'),
('FR15','Create a support request or complaint. CSR inspects the queue, responds and updates resolution; citizens view their own requests.','Citizen / CSR'),
('FR16','Review dashboards/performance and referred cases. Process biometric-complete applications as an issuance batch and export the register to CSV.','Director'),
('FR17','Create, edit and deactivate staff; maintain citizen accounts, service centres and fees through authorised administration routes.','Administrator'),
('FR18','Publish announcements/maintenance notices, and store in-app notifications for relevant application, payment, booking and staff events.','Admin / users'),
('FR19','Provide public guidance and a landing-page assistant. Public scope excludes private account access and state-changing actions.','Public visitor')],[.48,4.95,1.07],10)
h('Supporting requirements and data rules')
para('Email/SMS delivery, password recovery, file preview, QR generation, notification read state and shared navigation support the major functions. Common work was shared across the group. References, ownership links and workflow status must remain consistent between screens and persisted records.')
para('Appointments identify their service type and centre. Active booking counts include confirmed, checked-in and completed bookings; cancellation releases capacity. The server checks ownership and current state when a citizen changes a booking. Stale confirmations are rejected instead of overwriting a newer booking.')
h('Priority and completion boundaries')
para('The design backlog contains 31 stories: 20 high, 10 medium and one low priority [2]. Core portal workflows are implemented. Activity monitoring is represented by domain audit records and operational views; a general log browser is outside the delivered interface. These boundaries appear again in the implementation and future-work sections.')
page('Nonfunctional requirements and constraints')
table(['Quality','Required behaviour and implementation evidence','Verification boundary'],[
('Security','Use BCrypt password hashes [3]. Validate verification codes, role/owner access and tokens on protected state-changing forms.','Role/token tests exist; security certification is not claimed.'),
('Privacy','Restrict evidence access. Keep QR payloads to the appointment reference; avoid raw biometric captures and card credential storage.','Retention and deployment policies require further work.'),
('Integrity','Enforce keys, relationships and valid states. Use transactions/row locks for booking capacity and reject stale versions.','Check normal, full-slot, duplicate and stale-request paths.'),
('Usability','Use consistent navigation, labelled fields, inline feedback, a progress timeline and printable passes. Provide assisted visits.','Screens illustrate the UI; a formal user study is outstanding.'),
('Accessibility','Aim for keyboard access, clear labels, readable contrast and responsive layouts using WCAG 2.2 [4].','Conformance requires an accessibility audit.'),
('Reliability','Keep related writes transactional, explain failures and display maintenance notices when required.','Selected integration tests exist; uptime is unmeasured.'),
('Performance','Use server-rendered views and database lookups suitable for the classroom prototype.','Response-time, throughput and concurrency need load testing.'),
('Maintainability','Separate controller, service, repository, model and DTO work; use constructor injection and a Maven build.','Module tests and shared policies support regression checks.')],[1.02,3.53,1.95],10)
h('Constraints and assumptions')
para('The application is Java-based and developed within a semester [1]. The build declares Spring Boot 4.1.0 and Java 17; SQL Server and its schema must be configured before startup. Hibernate validates the schema rather than creating production tables automatically. Version-specific setup must be checked before demonstration.')
para('Email verification needs configured SMTP credentials and network access. SMS uses an HTTP phone gateway whose availability depends on the local network/device. Webcam scanning needs camera permission and a suitable browser origin. Credentials belong in environment configuration and must not appear in report screenshots.')
para('Appointment policies use Sri Lanka time. The simulated payment/capture stages demonstrate business flows but do not establish banking settlement, legal identity verification or actual biometric collection. An approved prototype application is not an official identity document.')
section_end(3)
h('Design',1)
h('Use cases for citizens and application officers')
para('Figures 1 and 2 refine interactions from the design document [2] to cover the final prototype. UML actors remain outside the system boundary and ovals represent user goals [5]. Authentication is a precondition for protected goals; registration and public walk-in services use their own verification steps.')
figure(TMP/'usecases-citizen-officer.png','Figure 1. Final citizen and application-officer use cases',width=6.3,height=6.4)
para('Citizens manage their own requests. NIC and passport officers use review and appointment operations within the document type permitted by their session. Approval enables the payment/appointment stages; corrections return responsibility to the applicant.',10.5)
page('Use cases for assisted service and administration')
para('The visitor route supports people who need an assisted visit. CSR, director and administrator responsibilities are separate: organisational seniority does not imply unrestricted technical access. The director also has CSR support access; administrator pages remain a separate authorised area.')
figure(TMP/'usecases-operations.png','Figure 2. Public visitor, CSR, director and administrator use cases',width=6.3,height=6.45)
para('Notifications support workflow events. The public assistant provides guidance without submitting applications or exposing accounts. Director printing records issuance and exports a register; it does not operate a physical printer.',10.5)
page('Database design for accounts and applications')
para('SQL Server stores the relational domain, while JPA maps entities/repositories to it [6]. Figure 3 presents the core account, identity and application structure, checked against schema constraints and JPA mappings. Shared keys separate common account/application data from citizen, staff, NIC and passport details.')
figure(TMP/'er-accounts.png','Figure 3. Account, application, evidence, review and delivery relationships',width=6.5,height=5.6)
table(['Relationship','Meaning and cardinality'],[
('UserAccount to CitizenAccount / StaffAccount','Each specialised row shares its user key; common login data is stored once.'),
('CitizenAccount to Application','One citizen can have many applications; each application belongs to one citizen.'),
('Application to NICApplication / PassportApplication','Document-specific detail uses the application key instead of duplicating the whole application.')],[2.35,4.15],9.5)
page('Database design for appointments and supporting records')
para('Figure 4 presents the appointment/support domain. Repeated account/application entities connect it to Figure 3. Primary keys identify records; foreign keys enforce ownership and dependencies. Operational settings for fees, announcements and maintenance complement this core model.')
figure(TMP/'er-operations.png','Figure 4. Appointment, slot, payment, support and notification relationships',width=6.5,height=6.0)
para('A centre owns many published slots and appointments. An appointment can link to its application, citizen and biometric slot; guest walk-ins use a separate service path. Check-in records identify the appointment and staff member. UserNotification joins users to notifications with per-user read state. Booking transactions lock relevant rows before checking capacity.',10.5)
page('Interface design for account access and progress')
para('The portal uses a shared government-style theme, consistent navigation and role-specific workbenches. These are sample-data prototype screenshots retained from the team draft; they demonstrate interface behaviour rather than live government operations.')
figure(media(7),'Figure 5. Unified citizen and staff login portal',width=5.5,height=2.8)
para('Login provides role selection and recovery navigation. Registration separates email verification, mobile verification and password creation so the next task is visible. The server validates the sequence and submitted values.',10.5)
figure(media(11),'Figure 6. Citizen application timeline and next action',width=6.2,height=2.8)
para('Tracking pairs the current state with completed and pending stages. A visible next action leads to payment or booking when eligible. Correction messages return the user to affected fields/evidence instead of requiring an unrelated new application.',10.5)
page('Interface design for booking and appointment passes')
figure(media(13),'Figure 7. Centre, date and time selection for biometric appointments',width=4.4,height=3.85)
para('Booking begins with the relevant application and centre. The calendar highlights available dates; published times expose the remaining choices. Server-side capacity checks remain authoritative when concurrent requests arrive.',10.5)
figure(media(12),'Figure 8. Confirmed appointment and printable QR pass',width=4.7,height=2.7)
para('The pass gives the visit information and appointment reference. Its QR carries the reference for authorised lookup; manual reference entry provides a centre fallback. Changes, printing and camera-based check-in have separate state/date checks.',10.5)
section_end(4)
h('Implementation',1)
h('Tools and technologies')
table(['Technology','Use in the prototype and selection rationale'],[
('Java 17 / Spring Boot 4.1.0 [7]','Typed domain/service code, dependency injection, request routing and transactions in one application.'),
('Thymeleaf [8], HTML/CSS/JavaScript','Server-rendered forms/workbenches, client feedback, calendar interaction and responsive navigation.'),
('Tailwind CSS [9] / shared theme','Utility styling with consistent portal colours, spacing and components.'),
('Spring Data JPA / Hibernate [10]','Repositories and entity mapping; schema validation exposes mismatches at startup.'),
('SQL Server / Microsoft JDBC [6]','Relational constraints, persistent workflow records, transactions and booking locks.'),
('Security Crypto / Bean Validation','BCrypt and validated forms; controller/session checks govern access.'),
('Spring Mail / HTTP SMS adapter','Email/mobile verification through configurable delivery services.'),
('ZXing 3.5.4 [11] / jsQR 1.4.0 [12]','QR pass generation and browser webcam scanning.'),
('Maven [13], JUnit 5 [14], MockMvc, Mockito [15]','Dependencies/builds and unit, controller and database-backed checks.'),
('Git / GitHub [16]','Member branches, integrated changes and shared source history.')],[1.95,4.55],9.7)
h('Architecture and design patterns')
para('LankaID is a layered monolith. A browser request reaches a Spring MVC controller, which checks session/role context and calls a service. Services hold business rules/transaction boundaries; repositories persist entities to SQL Server. Thymeleaf renders the model. OTP adapters, local file storage, QR generation and notification publishing connect to this flow.')
table(['Pattern or separation','Example and benefit'],[
('MVC / service layer','AppointmentController delegates to AppointmentService, separating presentation and business decisions.'),
('Repository','AppointmentRepository defines persistent lookup and locking without database access in templates.'),
('Observer','NotificationPublisher sends an event to InAppNotificationObserver, reducing module/storage coupling.'),
('Adapter / policy object','SmsSender isolates delivery; BookingChangePolicy centralises time/status rules for pages and server checks.')],[1.45,5.05],9.7)
page('Implemented review and decision workflow')
para('The main modules cover applications, verification, appointments, renewal/replacement/delivery and support. User, OTP, payment, notification and administration services connect them. The screens below illustrate how requirements become a connected workflow.')
figure(media(15),'Figure 9. NIC officer queue with submitted sample applications',width=6.0,height=2.8)
para('The queue presents applications relevant to the officer role. Review loads details and authorised evidence previews. Approval allows progression, correction returns work to the citizen, and rejection retains a reason. Director referral provides a recorded path for cases needing further review.',10.5)
figure(media(16),'Figure 10. Supporting evidence and officer decision controls',width=6.0,height=3.2)
para('The decision service records review details and updates the state. Notifications communicate the outcome; the tracking view exposes the next action. Reasons, current-state validation and ownership checks reduce ambiguous/repeated decisions. Automated tests exercise transitions and access rules.',10.5)
page('Implemented appointment and issuance workflow')
figure(media(17),'Figure 11. Schedule with booked counts and slot-capacity controls',width=6.0,height=3.2)
para('Officers publish capacity-limited slots. Booking/changes use transactional row locking and a booking version for consistent capacity and stale-request checks. BookingChangePolicy enforces the biometric 72-hour change window, missed-booking rules and check-in on the booked date.',10.5)
figure(media(18),'Figure 12. Biometric desk with QR scan and reference-number fallback',width=5.85,height=3.1)
para('After check-in, the officer records simulated biometric completion. The director selects biometric-complete applications, processes an issuance batch and exports the CSV register. PrintingExportService checks current state and calls the delivery service after issuance. Delivery/collection records the final handover; device capture and physical printing are outside the prototype.',10.5)
page('Testing evidence and implementation boundaries')
para('Tests use JUnit, MockMvc controller checks, mocked dependencies and transaction-backed integration workflows [14], [15]. The project includes tests for account/OTP rules, drafts/submission, decisions, payments, appointments, walk-ins and delivery. These checks demonstrate selected behaviour, not a measured percentage of code coverage.')
table(['Area','Example automated evidence','Behaviour checked'],[
('Applications','CitizenApplicationWorkflowIntegrationTests','Draft, evidence and submission transitions.'),
('Review','VerificationDecisionWorkflowIntegrationTests','Approval, rejection, correction and decisions.'),
('Bookings','BookingChangePolicyTests; BookingChangeWorkflowIntegrationTests','Time/status eligibility, stale changes and workflow.'),
('Payment','PaymentBiometricWorkflowIntegrationTests','Payment prerequisites and appointment progression.'),
('Walk-in / OTP','WalkInAppointmentWorkflowIntegrationTests; SmsOtpServiceTests','Verified booking, expiry and attempt rules.'),
('Routes / schema','UiRouteSmokeTests; UiAndSchemaConsistencyTests','Role-sensitive routes and view/schema consistency.')],[1.1,3.0,2.4],9.5)
h('Recorded results and remaining findings')
para('Stored Surefire XML reports contain 295 executed cases across 63 report files: 293 passed, two failed, zero errors and zero skipped. These historical files can span different runs; they do not establish a fresh full-suite result against the final submission.')
para('Both failures are in MarickarVerificationTestCases. One expects HTTP 403 for an unauthorised CSR request but records HTTP 404; the other expects an official-notification email call that was not observed. Expected behaviour must be reconciled with current routes/notification integration, followed by a clean regression run before claiming that every test passes.')
h('Ethical and operational considerations')
para('Identity evidence is sensitive. Access checks, limited QR payloads, hashes and traceable decisions support privacy/accountability. A human officer makes decisions; the assistant provides public guidance. Assisted visits reduce exclusion. Deployment needs a retention policy, encrypted storage and formal security/accessibility/performance checks. OWASP ASVS provides a security-review reference [17].')
section_end(5)
h('Project Management',1)
h('Agile approach and sprint summary')
para('The design document uses a Scrum-oriented backlog and five planned sprints [2], [18]. It identifies product-owner/Scrum-master responsibilities conceptually as the Identity Service Director and IT project lead. These are planning roles, not evidence that a government director joined the student project. Development, interface review and testing were shared by five students.')
table(['Sprint plan','Goal and planned stories','Evidence in the final prototype'],[
('1 / weeks 1-3 / 81 h','Accounts, administration, roles, OTP and activity monitoring.','Registration/login, administration and domain audit records; no general log browser.'),
('2 / weeks 4-6 / 51 h','NIC/passport applications, uploads and payment.','Drafts, forms, evidence and simulated payment.'),
('3 / weeks 7-9 / 103 h','Officer review, verify, approve/reject and notify.','Separate queues, recorded reviews, correction/referral and notifications.'),
('4 / weeks 10-12 / 33 h','Capacity slots, booking, QR passes and check-in.','Slot/calendar functions, QR/manual check-in and simulated completion.'),
('5 / weeks 13-15 / 110 h','Tracking, support, maintenance, renewals, replacement, delivery and booking changes.','Tracking, assisted visits, support, renewal/replacement, handover and changes.')],[1.35,2.55,2.6],10)
para('The 378 hours are estimates from the design plan, not measured effort. That plan puts Sprint 5 in weeks 13-15, while final evaluation is scheduled in week 13 [1]. Implementation therefore overlaps the original sequence; dated milestones on the next page provide the more reliable picture of completed development.')
h('Coordination and integration')
para('Each primary function had an owner; common services required cross-module coordination. Member branches/merges maintained shared source history. DTOs, status rules, navigation and a common theme aligned citizen/staff workflows. Database scripts addressed changes that otherwise prevented teammates running the same version.')
para('A practical completion criterion is a working user flow with server validation, persisted state, understandable feedback and regression checks. Recording review/retrospective outcomes alongside the backlog and requiring a clean shared build before merging would strengthen this approach.')
page('Task distribution and development milestones')
table(['Member and ID','Primary responsibility','Connected/shared contribution'],[
('Weerakoon W M Y B\nIT25100565','Biometric/appointment management.','Slots, changes, QR passes/check-in, officer administration and shared theme.'),
('Dharshan S\nIT25100604','Citizen support/service management.','Walk-ins, mobile verification, CSR queue and assisted service/handover.'),
('Jeniviya G A\nIT25100689','Citizen application management.','NIC/passport forms, drafts, uploads, corrections and tracking.'),
('Rithika R K\nIT25100885','Renewal/replacement/delivery.','Existing-document requests, delivery/collection and receipt/handover.'),
('Marickar A A A\nIT25101235','User management and review/approval.','Accounts, OTP/payment integration, officer queues and integration support.')],[1.62,2.16,2.72],10)
para('The assignment describes six-member groups; the supplied cover/design list five students. The report retains those five and records shared tasks. Primary ownership does not imply that every file in a module was written by one person.',10.5)
h('Dated milestones')
table(['Period in 2026','Development focus'],[
('10-18 August','Repository foundations, initial templates/tests and module structure.'),
('20 August-10 September','Applications, uploads/drafts, review queues, payment and initial renewal/walk-in flows.'),
('14-19 September','Email/SMS, delivery/collection, booking changes, staff views and schema work.'),
('27 September','Dated design document: sprints, UML diagrams and ethics [2].'),
('5-10 October','UI integration, guidance, notices, maintenance, booking safeguards, tests and documentation.')],[1.48,5.02],10)
para('Milestones summarise dated source history and the design-document date. They are development phases, not meeting minutes or a claim that sprint weeks were followed exactly. Regression evidence, repository access and report review form part of submission preparation.')
section_end(6)
h('Conclusion and Future Work',1)
h('Summary of achievements')
para('LankaID connects the main stages of an identity-service request in one academic prototype. Citizens have verified registration, drafts/evidence, tracked reviews, simulated payment, appointment passes and delivery/collection records. Officers have document-specific queues and centre operations; CSR provides assisted service; the director manages referrals/issuance records; the administrator maintains operating settings.')
para('The main outcome is a connected workflow with explicit responsibilities and controlled states. UML, relational design, layered Java code and automated checks link software-engineering activities to delivered screens. Live payments, government-data integration, biometric devices and physical production remain outside the prototype scope.')
h('Challenges and responses')
table(['Challenge','Response and remaining lesson'],[
('Database drift','Schema validation exposed mismatches. Shared schema/migration scripts improved setup; versioned migrations would improve upgrades.'),
('Concurrent/stale booking','Row locks, booking versions and a shared policy protect capacity/changes. Boundary cases must stay in regression checks.'),
('Module integration','Shared statuses, navigation and theme connected independently developed screens. Earlier contracts would reduce rework.'),
('External dependencies','Gateway interfaces and mocks limit coupling. Demonstration still requires verified SMTP/SMS configuration.'),
('Regression evidence','Historical results record two failed assertions. A clean dated rerun is needed before claiming a fully passing final suite.')],[1.55,4.95],10)
h('Prioritised future work')
para('First, resolve regression findings, standardise migrations and complete a security review with framework-level CSRF protection and encrypted evidence storage. Document retention/consent procedures and verify every role/owner boundary before deployment.')
para('Next, integrate certified payment, authorised government-data interfaces, biometric devices and delivery/printing services with recovery/audit checks. Then add Sinhala/Tamil, perform accessibility/usability studies, introduce reminders and load-test paginated staff lists. Measured outcomes should guide scaling decisions.')
section_end(7)
h('Individual Contribution Teamwork and Lessons Learned',1)
para('The table summarises each member’s primary work and the engineering lessons associated with the implemented modules. Shared responsibilities include integration, interface consistency and regression checking.')
table(['Member','Role and contribution','Challenge and response','Lessons and reflection'],[
('Weerakoon W M Y B\nIT25100565','Biometrics/appointments: slot publication, booking calendar, QR passes, check-in and change policy. Also officer administration and shared interface styling.','Avoiding double booking and outdated changes required row locks, a booking version and shared date/status rules. Workflow tests exercise the checks; schema/setup guidance helps align environments.','Concurrency/time rules need explicit tests and one authoritative definition. A consistent UI helped integration. Earlier migrations and boundary tests would reduce late hardening.'),
('Dharshan S\nIT25100604','Citizen support and assisted service: walk-ins, mobile verification, CSR queues/responses and service/handover assistance.','Guest visits require verified contact without assuming a citizen account. The flow separates OTP, booking/pass recovery and CSR check-in. Gateway delivery is isolated for mocked testing.','Assisted users need recovery as well as a successful booking. Separate guest/account ownership makes the flow clearer. Earlier handover checks would reduce integration surprises.'),
('Jeniviya G A\nIT25100689','Citizen applications: NIC/passport forms, drafts, evidence, correction/resubmission and progress views.','Multi-step requests must keep draft data/files consistent and prevent invalid submission. Service validation, ownership, file-size limits and explicit correction/submission states address this.','Client feedback helps, but server validation remains authoritative. Full-journey tests expose issues that isolated forms miss. Earlier shared statuses/evidence rules would reduce rework.')],[1.03,1.67,2.05,1.75],9.5)
page('Individual reflections and team learning')
table(['Member','Role and contribution','Challenge and response','Lessons and reflection'],[
('Rithika R K\nIT25100885','Renewal/replacement/delivery: existing-document requests, evidence, delivery/collection choice and receipt/handover.','The module depends on application/payment/issuance states from other modules. Identity linking and eligibility separate renewal from new requests. Updates after printing connect delivery to the same application.','Clear inputs/transitions are essential before integration. Delivery belongs to the same lifecycle. Earlier approval-to-collection tests would improve confidence.'),
('Marickar A A A\nIT25101235','User access and review/approval: registration/login, contact verification, restricted queues, evidence and decisions. Supported OTP/payment and integration.','Sensitive routes require role checks and predictable outcomes. Sessions, BCrypt and recorded reasons support this. Historical tests reveal access-response and email expectations requiring reconciliation.','Authentication and authorisation need separate evaluation. A decision screen does not prove side effects/forbidden paths. Earlier role and notification tests would improve integration.')],[1.03,1.67,2.05,1.75],9.5)
h('Teamwork and improvement')
para('The strongest shared result was the module connection: citizen submission feeds review, approval feeds payment/booking, biometric completion feeds issuance, and issuance feeds delivery. Agreement on records, status names and transition ownership made those handoffs possible. Shared navigation/theme made them easier to follow.')
para('Coordination can improve through versioned migrations, common response/notification contracts and regression results regenerated for the exact submission. Planned effort should remain separate from completion dates so schedule problems stay visible.')
para('Future work should pair a named feature owner with an agreed acceptance flow and another member checking the handoff. Regular integration checks are easier to diagnose than a late combined failure. Personal module knowledge and understanding of the full lifecycle are both necessary for the demonstration and viva.')
section_end(8)
h('References',1)
para('IEEE reference numbers follow first use. Online resources were accessed on 13 October 2026. Course/design documents and the repository provide project evidence; external documentation supports modelling and implementation techniques.')
refs=[
'[1] Sri Lanka Institute of Information Technology, “SE2030 Software Engineering Group Assignment Specification,” Faculty of Computing, 2026, unpublished course document.',
'[2] Group 2026-Y2-S1-MTR-02, “Web Based Identity Issuing Service Continuous Assessment Design Document,” Faculty of Computing, SLIIT, Sep. 27, 2026, unpublished project document.',
'[3] Spring Security, “Password Storage,” Spring Security Reference. Accessed: Oct. 13, 2026. [Online]. Available: https://docs.spring.io/spring-security/reference/features/authentication/password-storage.html',
'[4] World Wide Web Consortium, “Web Content Accessibility Guidelines (WCAG) 2.2,” W3C Recommendation. Accessed: Oct. 13, 2026. [Online]. Available: https://www.w3.org/TR/WCAG22/',
'[5] Object Management Group, “Unified Modeling Language Specification,” ver. 2.5.1, Dec. 2017. Accessed: Oct. 13, 2026. [Online]. Available: https://www.omg.org/spec/UML/2.5.1/',
'[6] Microsoft, “SQL Server Technical Documentation,” Microsoft Learn. Accessed: Oct. 13, 2026. [Online]. Available: https://learn.microsoft.com/en-us/sql/sql-server/',
'[7] Spring, “Spring Boot Reference Documentation.” Accessed: Oct. 13, 2026. [Online]. Available: https://docs.spring.io/spring-boot/index.html',
'[8] The Thymeleaf Team, “Thymeleaf Documentation.” Accessed: Oct. 13, 2026. [Online]. Available: https://www.thymeleaf.org/documentation.html',
'[9] Tailwind Labs, “Tailwind CSS Documentation.” Accessed: Oct. 13, 2026. [Online]. Available: https://tailwindcss.com/docs'
]
for s in refs:para(s,10.5)
page('References continued')
refs=[
'[10] Hibernate, “Hibernate ORM Documentation.” Accessed: Oct. 13, 2026. [Online]. Available: https://hibernate.org/orm/documentation/',
'[11] ZXing Contributors, “ZXing Barcode Scanning Library,” GitHub. Accessed: Oct. 13, 2026. [Online]. Available: https://github.com/zxing/zxing',
'[12] jsQR Contributors, “jsQR JavaScript QR Code Reading Library,” GitHub. Accessed: Oct. 13, 2026. [Online]. Available: https://github.com/cozmo/jsQR',
'[13] Apache Software Foundation, “Apache Maven.” Accessed: Oct. 13, 2026. [Online]. Available: https://maven.apache.org/',
'[14] JUnit Team, “JUnit 5 User Guide,” ver. 5.14.0. Accessed: Oct. 13, 2026. [Online]. Available: https://docs.junit.org/5.14.0/user-guide/',
'[15] Mockito Contributors, “Mockito Framework.” Accessed: Oct. 13, 2026. [Online]. Available: https://site.mockito.org/',
'[16] Group 2026-Y2-S1-MTR-02, “LankaID Identity Issuing Web Service Application,” source repository. Accessed locally: Oct. 13, 2026. [Online]. Available: https://github.com/IT25101235/LankaID',
'[17] OWASP Foundation, “Application Security Verification Standard.” Accessed: Oct. 13, 2026. [Online]. Available: https://owasp.org/projects/asvs',
'[18] K. Schwaber and J. Sutherland, “The Scrum Guide,” Nov. 2020. Accessed: Oct. 13, 2026. [Online]. Available: https://scrumguides.org/scrum-guide.html'
]
for s in refs:para(s,10.5)
section_end(9)
h('Appendix',1)
h('Source and reproducibility')
para('Source, database scripts and setup guidance are at https://github.com/IT25101235/LankaID [16]. Configure SQL Server, apply the schema/migrations and keep mail/SMS credentials outside source. The Maven wrapper builds/runs the application and generates Surefire results. Demonstrate with sample records and confirm evaluator repository access.')
h('Requirements to evidence map')
table(['Requirement group','Implementation evidence','Report evidence'],[
('FR01-FR02 / accounts','User_Mgmt; OTP_Mgmt; registration/session tests.','Account screen, constraints and testing.'),
('FR03-FR06 / applications','CitizenApplicationService; identity linking; renewal/replacement module.','Use cases, ER model, tracking and application tests.'),
('FR07-FR10 / payment, booking, handover','PaymentService; AppointmentService; BookingChangePolicy; DocumentDeliveryService.','Booking/pass views, operations and simulations.'),
('FR11-FR13 / review and centre work','VerificationReviewService; BiometricScheduleService; BiometricService.','Queue, decisions, schedule and check-in images.'),
('FR14-FR15 / assisted service','WalkInAppointmentService; CitizenSupportService.','Public/CSR use cases and walk-in/OTP tests.'),
('FR16-FR19 / operations','PrintingExportService; Admin_Mgmt; Notification_Mgmt; Chatbot_Mgmt.','Role model, issuance, oversight and public scope.')],[1.55,2.65,2.3],10)
h('Traceability of project evidence')
para('Requirements express the final prototype scope. Sprint goals/effort estimates originate in the dated design [2]; milestones summarise development history. UI assets come from the team draft; ER diagrams reflect checked keys and mappings. Stored XML supports recorded test results; the submission version should be rerun to replace historical evidence.')
uf=doc.settings._element.find(qn('w:updateFields'))
if uf is None:uf=OxmlElement('w:updateFields');doc.settings._element.append(uf)
uf.set(qn('w:val'),'true')
doc.core_properties.title='Web Based Identity Issuing Service Final Report';doc.core_properties.subject='SE2030 final project report group 2026-Y2-S1-MTR-02';doc.core_properties.author='Group 2026-Y2-S1-MTR-02'
doc.save(FINAL)
editable={'word/document.xml','word/styles.xml','word/settings.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','docProps/app.xml'}
with zipfile.ZipFile(SOURCE) as baseline,zipfile.ZipFile(FINAL) as current:
 parts={n:current.read(n) for n in current.namelist()}
 for n in baseline.namelist():
  if n not in editable:parts[n]=baseline.read(n)
with zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
print('Saved',FINAL)
print('Sections',len(Document(FINAL).sections),'tables',len(Document(FINAL).tables),'figures',len(Document(FINAL).inline_shapes))
print('Source hash',hashlib.sha256(SOURCE.read_bytes()).hexdigest())
